import pandas as pd
import openpyxl
from openpyxl import load_workbook

# Verificar os dados na aba Dados_NFSe
print("Verificando a aba Dados_NFSe:")
df = pd.read_excel('output/teste_iss_final.xlsx', sheet_name='Dados_NFSe')

# Verificar se o Total ISS está sendo calculado corretamente
df['Total ISS Calculado'] = df['ISS Próprio (R$)'] + df['ISS Retido (R$)']
df['Diferença Total ISS'] = df['Total ISS (R$)'] - df['Total ISS Calculado']

# Verificar se há diferenças no Total ISS
diff_total_iss = df[abs(df['Diferença Total ISS']) > 0.01]
if not diff_total_iss.empty:
    print(f"\nEncontradas {len(diff_total_iss)} linhas com diferença no Total ISS")
    print(diff_total_iss[['Número da Nota', 'ISS Próprio (R$)', 'ISS Retido (R$)', 'Total ISS (R$)', 'Total ISS Calculado', 'Diferença Total ISS']].head())
else:
    print("\nNenhuma diferença encontrada no Total ISS na aba Dados_NFSe.")

# Verificar a aba Resumo_NFSe
print("\nVerificando a aba Resumo_NFSe:")
wb = load_workbook('output/teste_iss_final.xlsx')
ws = wb['Resumo_NFSe']

# Encontrar todas as tabelas na aba Resumo_NFSe
tabelas = []
for row in range(1, ws.max_row + 1):
    cell_value = ws.cell(row=row, column=1).value
    if cell_value and isinstance(cell_value, str) and cell_value.upper() == "SITUAÇÃO":
        tabelas.append(("Resumo por Situação", row))
    elif cell_value and isinstance(cell_value, str) and cell_value.upper() == "COMPETÊNCIA":
        tabelas.append(("Resumo por Competência", row))
    elif cell_value and isinstance(cell_value, str) and cell_value.upper() == "CÓDIGO DO SERVIÇO":
        tabelas.append(("Resumo por Código do Serviço", row))

# Verificar cada tabela
for tabela_nome, header_row in tabelas:
    print(f"\nVerificando tabela: {tabela_nome} (linha {header_row})")
    
    # Encontrar as colunas de ISS próprio, ISS retido e Total ISS
    iss_proprio_col = None
    iss_retido_col = None
    total_iss_col = None
    
    for col in range(1, ws.max_column + 1):
        cell_value = ws.cell(row=header_row, column=col).value
        if cell_value and "ISS PRÓPRIO" in str(cell_value).upper():
            iss_proprio_col = col
        elif cell_value and "ISS RETIDO" in str(cell_value).upper():
            iss_retido_col = col
        elif cell_value and "TOTAL ISS" in str(cell_value).upper():
            total_iss_col = col
    
    if iss_proprio_col and iss_retido_col:
        print(f"Coluna de ISS Próprio encontrada na coluna {iss_proprio_col}")
        print(f"Coluna de ISS Retido encontrada na coluna {iss_retido_col}")
        if total_iss_col:
            print(f"Coluna de Total ISS encontrada na coluna {total_iss_col}")
        else:
            print("Coluna de Total ISS NÃO encontrada!")
        
        # Verificar os valores
        problemas = 0
        for row in range(header_row + 1, min(header_row + 20, ws.max_row + 1)):
            # Verificar se a linha não está vazia
            if ws.cell(row=row, column=1).value is None:
                continue
                
            iss_proprio = ws.cell(row=row, column=iss_proprio_col).value
            iss_retido = ws.cell(row=row, column=iss_retido_col).value
            
            if iss_proprio is not None and iss_retido is not None:
                total_calculado = iss_proprio + iss_retido
                
                if total_iss_col:
                    total_iss = ws.cell(row=row, column=total_iss_col).value
                    if total_iss is not None and abs(total_iss - total_calculado) > 0.01:
                        problemas += 1
                        print(f"Problema na linha {row}: ISS Próprio={iss_proprio}, ISS Retido={iss_retido}, Total ISS={total_iss}, Total Calculado={total_calculado}")
                else:
                    problemas += 1
                    print(f"Problema na linha {row}: Coluna Total ISS não existe")
        
        if problemas == 0:
            print(f"Nenhum problema encontrado na tabela {tabela_nome}")
        else:
            print(f"Encontrados {problemas} problemas na tabela {tabela_nome}")
    else:
        print(f"Colunas de ISS não encontradas na tabela {tabela_nome}")
