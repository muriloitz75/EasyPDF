import pandas as pd
import openpyxl
from openpyxl import load_workbook

# Carregar o arquivo Excel original (com todas as notas)
df_original = pd.read_excel('output/teste_validadas.xlsx', sheet_name='Dados_NFSe')

# Contar o número de notas canceladas
notas_canceladas = df_original[df_original['Situação'] == 'CANCELADA']
print(f"Total de notas canceladas no arquivo original: {len(notas_canceladas)}")

# Verificar os totais na aba Dados_NFSe
print("\nTotais na aba Dados_NFSe:")
print(f"Total de notas: {len(df_original)}")
print(f"Total de notas não canceladas: {len(df_original) - len(notas_canceladas)}")

# Carregar o arquivo Excel novo (sem notas canceladas nos cálculos)
df_novo = pd.read_excel('output/teste_sem_canceladas.xlsx', sheet_name='Dados_NFSe')

# Verificar se o número total de notas é o mesmo (incluindo canceladas)
print(f"\nTotal de notas no novo arquivo: {len(df_novo)}")

# Carregar a aba Resumo_NFSe do arquivo original
wb_original = load_workbook('output/teste_validadas.xlsx')
ws_original = wb_original['Resumo_NFSe']

# Carregar a aba Resumo_NFSe do novo arquivo
wb_novo = load_workbook('output/teste_sem_canceladas.xlsx')
ws_novo = wb_novo['Resumo_NFSe']

# Função para encontrar a linha de total na tabela de competência
def encontrar_linha_total_competencia(ws):
    for row in range(1, ws.max_row + 1):
        if ws.cell(row=row, column=1).value == 'TOTAL GERAL':
            return row
    return None

# Função para encontrar a linha de total na tabela de código de serviço
def encontrar_linha_total_codigo_servico(ws):
    for row in range(1, ws.max_row + 1):
        if ws.cell(row=row, column=1).value == 'TOTAL':
            return row
    return None

# Verificar a tabela de resumo por competência
linha_total_competencia_original = encontrar_linha_total_competencia(ws_original)
linha_total_competencia_novo = encontrar_linha_total_competencia(ws_novo)

if linha_total_competencia_original and linha_total_competencia_novo:
    # Obter os valores de quantidade de notas
    qtd_notas_original = ws_original.cell(row=linha_total_competencia_original, column=2).value
    qtd_notas_novo = ws_novo.cell(row=linha_total_competencia_novo, column=2).value
    
    print("\nTabela de Resumo por Competência:")
    print(f"Quantidade de notas no arquivo original: {qtd_notas_original}")
    print(f"Quantidade de notas no novo arquivo (sem canceladas): {qtd_notas_novo}")
    print(f"Diferença: {qtd_notas_original - qtd_notas_novo}")
    
    # Verificar se a diferença corresponde ao número de notas canceladas
    if qtd_notas_original - qtd_notas_novo == len(notas_canceladas):
        print("✓ As notas canceladas foram excluídas corretamente da tabela de Resumo por Competência!")
    else:
        print("✗ A diferença não corresponde ao número de notas canceladas!")

# Verificar a tabela de resumo por código de serviço
linha_total_codigo_servico_original = encontrar_linha_total_codigo_servico(ws_original)
linha_total_codigo_servico_novo = encontrar_linha_total_codigo_servico(ws_novo)

if linha_total_codigo_servico_original and linha_total_codigo_servico_novo:
    # Obter os valores de quantidade de notas
    qtd_notas_original = ws_original.cell(row=linha_total_codigo_servico_original, column=2).value
    qtd_notas_novo = ws_novo.cell(row=linha_total_codigo_servico_novo, column=2).value
    
    print("\nTabela de Resumo por Código de Serviço:")
    print(f"Quantidade de notas no arquivo original: {qtd_notas_original}")
    print(f"Quantidade de notas no novo arquivo (sem canceladas): {qtd_notas_novo}")
    print(f"Diferença: {qtd_notas_original - qtd_notas_novo}")
    
    # Verificar se a diferença corresponde ao número de notas canceladas
    if qtd_notas_original - qtd_notas_novo == len(notas_canceladas):
        print("✓ As notas canceladas foram excluídas corretamente da tabela de Resumo por Código de Serviço!")
    else:
        print("✗ A diferença não corresponde ao número de notas canceladas!")
