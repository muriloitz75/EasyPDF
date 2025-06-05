import pandas as pd
import openpyxl
from openpyxl import load_workbook

# Carregar o arquivo Excel
df = pd.read_excel('output/teste_sem_canceladas.xlsx', sheet_name='Dados_NFSe')

# Contar o número de notas por situação
contagem_situacao = df['Situação'].value_counts()
print("Contagem de notas por situação:")
print(contagem_situacao)

# Contar o número de notas canceladas
notas_canceladas = df[df['Situação'] == 'CANCELADA']
print(f"\nTotal de notas canceladas: {len(notas_canceladas)}")

# Verificar os totais na aba Dados_NFSe
print("\nTotais na aba Dados_NFSe:")
print(f"Total de notas: {len(df)}")
print(f"Total de notas não canceladas: {len(df) - len(notas_canceladas)}")

# Carregar a aba Resumo_NFSe
wb = load_workbook('output/teste_sem_canceladas.xlsx')
ws = wb['Resumo_NFSe']

# Função para encontrar a linha de total na tabela de competência
def encontrar_linha_total_competencia(ws):
    for row in range(1, ws.max_row + 1):
        if ws.cell(row=row, column=1).value == 'TOTAL GERAL':
            return row
    return None

# Função para encontrar a linha de total na tabela de código de serviço
def encontrar_linha_total_codigo_servico(ws):
    for row in range(1, ws.max_row + 1):
        if ws.cell(row=row, column=1).value == 'TOTAL':
            return row
    return None

# Verificar a tabela de resumo por competência
linha_total_competencia = encontrar_linha_total_competencia(ws)

if linha_total_competencia:
    # Obter os valores de quantidade de notas
    qtd_notas = ws.cell(row=linha_total_competencia, column=2).value
    
    print("\nTabela de Resumo por Competência:")
    print(f"Quantidade de notas: {qtd_notas}")
    
    # Verificar se a quantidade corresponde ao número de notas não canceladas
    if qtd_notas == len(df) - len(notas_canceladas):
        print("✓ As notas canceladas foram excluídas corretamente da tabela de Resumo por Competência!")
    else:
        print(f"✗ A quantidade de notas na tabela ({qtd_notas}) não corresponde ao número de notas não canceladas ({len(df) - len(notas_canceladas)})!")

# Verificar a tabela de resumo por código de serviço
linha_total_codigo_servico = encontrar_linha_total_codigo_servico(ws)

if linha_total_codigo_servico:
    # Obter os valores de quantidade de notas
    qtd_notas = ws.cell(row=linha_total_codigo_servico, column=2).value
    
    print("\nTabela de Resumo por Código de Serviço:")
    print(f"Quantidade de notas: {qtd_notas}")
    
    # Verificar se a quantidade corresponde ao número de notas não canceladas
    if qtd_notas == len(df) - len(notas_canceladas):
        print("✓ As notas canceladas foram excluídas corretamente da tabela de Resumo por Código de Serviço!")
    else:
        print(f"✗ A quantidade de notas na tabela ({qtd_notas}) não corresponde ao número de notas não canceladas ({len(df) - len(notas_canceladas)})!")
